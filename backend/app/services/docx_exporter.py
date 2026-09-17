from pathlib import Path
from typing import Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from app.core.config import settings

from html.parser import HTMLParser
import html

def set_cell_border(cell, **kwargs):
    """
    Set cell borders in Word table: top, bottom, left, right.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>\n'
                          f'<w:top w:val="{kwargs.get("top", "none")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("top_color", "auto")}"/>\n'
                          f'<w:left w:val="{kwargs.get("left", "none")}" w:sz="{kwargs.get("left_sz", "4")}" w:space="0" w:color="{kwargs.get("left_color", "auto")}"/>\n'
                          f'<w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("bottom_color", "auto")}"/>\n'
                          f'<w:right w:val="{kwargs.get("right", "none")}" w:sz="{kwargs.get("right_sz", "4")}" w:space="0" w:color="{kwargs.get("right_color", "auto")}"/>\n'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def set_cell_background(cell, hex_color: str):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

class DocxHtmlParser(HTMLParser):
    def __init__(
        self,
        paragraph,
        default_font: str = "Arial",
        default_size: Pt = Pt(10),
        default_bold: bool = False,
        default_italic: bool = False,
        default_underline: bool = False,
        default_color: Optional[RGBColor] = None
    ):
        super().__init__()
        self.paragraph = paragraph
        self.default_font = default_font
        self.default_size = default_size
        self.default_color = default_color
        
        self.bold = 1 if default_bold else 0
        self.italic = 1 if default_italic else 0
        self.underline = 1 if default_underline else 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ("b", "strong"):
            self.bold += 1
        elif tag in ("i", "em"):
            self.italic += 1
        elif tag in ("u", "ins"):
            self.underline += 1
        elif tag in ("br", "br/"):
            self.paragraph.add_run().add_break()
        elif tag == "li":
            run = self.paragraph.add_run("• ")
            run.font.name = self.default_font
            run.font.size = self.default_size

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("b", "strong"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("i", "em"):
            self.italic = max(0, self.italic - 1)
        elif tag in ("u", "ins"):
            self.underline = max(0, self.underline - 1)
        elif tag == "li":
            self.paragraph.add_run().add_break()

    def handle_data(self, data):
        if not data:
            return
        lines = data.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                self.paragraph.add_run().add_break()
            if line:
                run = self.paragraph.add_run(line)
                run.font.name = self.default_font
                run.font.size = self.default_size
                run.bold = bool(self.bold > 0)
                run.italic = bool(self.italic > 0)
                run.underline = bool(self.underline > 0)
                if self.default_color:
                    run.font.color.rgb = self.default_color

def add_html_formatted_text(
    paragraph,
    text: str,
    default_font: str = "Arial",
    default_size: Pt = Pt(10),
    default_bold: bool = False,
    default_italic: bool = False,
    default_underline: bool = False,
    default_color: Optional[RGBColor] = None
):
    if not text:
        return
    if "<" not in str(text):
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                paragraph.add_run().add_break()
            if line:
                run = paragraph.add_run(line)
                run.font.name = default_font
                run.font.size = default_size
                run.bold = default_bold
                run.italic = default_italic
                run.underline = default_underline
                if default_color:
                    run.font.color.rgb = default_color
        return

    parser = DocxHtmlParser(
        paragraph=paragraph,
        default_font=default_font,
        default_size=default_size,
        default_bold=default_bold,
        default_italic=default_italic,
        default_underline=default_underline,
        default_color=default_color
    )
    parser.feed(str(text))

def export_paper_to_docx(
    paper_data: Dict[str, Any],
    output_path: Path,
    include_answer_key: bool = False,
    answer_key_data: Optional[Dict[str, Any]] = None
) -> Path:
    doc = docx.Document()
    
    # Page setup: A4 Portrait, 10mm margins (approx 0.39 inches)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    
    metadata = paper_data.get("metadata", {})
    school_name = metadata.get("school_name", "TRINITY HIGH SCHOOL & JUNIOR COLLEGE")
    tagline = metadata.get("tagline", "KNOWLEDGE IS WISDOM")
    exam_title = metadata.get("exam_title", "प्रथम घटक चाचणी")
    cls_name = metadata.get("class_name") or metadata.get("class") or "10"
    subject = metadata.get("subject", "हिंदी (लोकभारती)")
    duration = metadata.get("duration", "२ घंटे")
    total_marks = paper_data.get("total_marks", metadata.get("total_marks", 40))
    has_logo = metadata.get("has_logo", True)
    
    # Header Table: Left Logo, Right Center School Info
    header_table = doc.add_table(rows=1, cols=2 if has_logo else 1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    logo_path = settings.ASSETS_DIR / "1000358221.png"
    if not logo_path.exists():
        src_path = Path(r"C:\Users\User\Hindi paper\1000358221.png")
        if src_path.exists():
            logo_path = src_path

    if has_logo and logo_path.exists():
        cell_logo = header_table.cell(0, 0)
        cell_logo.width = Inches(1.2)
        p_logo = cell_logo.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_logo = p_logo.add_run()
        run_logo.add_picture(str(logo_path), width=Inches(1.1))
        
        cell_info = header_table.cell(0, 1)
        cell_info.width = Inches(6.2)
    else:
        cell_info = header_table.cell(0, 0)
        cell_info.width = Inches(7.4)
        
    p_info = cell_info.paragraphs[0]
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sch = p_info.add_run(f"{school_name}\n")
    run_sch.font.name = "Times New Roman"
    run_sch.font.size = Pt(17)
    run_sch.font.bold = True
    
    run_tag = p_info.add_run(f"{tagline}\n")
    run_tag.font.name = "Arial"
    run_tag.font.size = Pt(9.5)
    run_tag.font.italic = True
    run_tag.font.bold = True
    
    run_ex = p_info.add_run(f"{exam_title}")
    run_ex.font.name = "Arial"
    run_ex.font.size = Pt(13)
    run_ex.font.bold = True
    
    # Bottom border line under header
    for cell in header_table.rows[0].cells:
        set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

    # Details Table: 4 columns
    doc.add_paragraph()  # spacing
    details_table = doc.add_table(rows=1, cols=4)
    details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details_table.autofit = False
    
    widths = [Inches(2.0), Inches(2.2), Inches(1.6), Inches(1.6)]
    col_texts = [
        f"कक्षा: {cls_name}वीं",
        f"विषय: {subject}",
        f"समय: {duration}",
        f"कुल अंक: {total_marks}"
    ]
    aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    
    for i in range(4):
        cell = details_table.cell(0, i)
        cell.width = widths[i]
        set_cell_border(cell, bottom="single", bottom_sz="8", bottom_color="000000")
        p = cell.paragraphs[0]
        p.alignment = aligns[i]
        run = p.add_run(col_texts[i])
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.bold = True
        
    # Instructions
    instructions = paper_data.get("general_instructions", [])
    if instructions:
        p_inst = doc.add_paragraph()
        p_inst.paragraph_format.space_before = Pt(6)
        p_inst.paragraph_format.space_after = Pt(4)
        run_lbl = p_inst.add_run("सूचनाएँ : ")
        run_lbl.font.bold = True
        run_lbl.font.size = Pt(9.5)
        for inst in instructions:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Inches(0.2)
            p_item.paragraph_format.space_before = Pt(0)
            p_item.paragraph_format.space_after = Pt(2)
            add_html_formatted_text(p_item, f"• {inst}", default_font="Arial", default_size=Pt(9))

    # Vibhags & Questions
    sections = paper_data.get("sections", [])
    for sec in sections:
        sec_title = sec.get("section_title", "विभाग")
        sec_marks = sec.get("section_marks", 0)
        
        # Vibhag heading banner table
        sec_table = doc.add_table(rows=1, cols=1)
        sec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sec_cell = sec_table.cell(0, 0)
        sec_cell.width = Inches(7.4)
        set_cell_background(sec_cell, "F0F0F0")
        set_cell_border(sec_cell, top="single", bottom="single", left="single", right="single",
                        top_sz="6", bottom_sz="6", left_sz="6", right_sz="6")
        
        p_sec = sec_cell.paragraphs[0]
        p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sec = p_sec.add_run(f"{sec_title} ({sec_marks} अंक)")
        run_sec.font.name = "Arial"
        run_sec.font.size = Pt(11)
        run_sec.font.bold = True
        
        # Questions in section
        for q in sec.get("questions", []):
            q_num = q.get("question_number", "")
            q_text = q.get("question_text", "")
            q_marks = q.get("marks", 0)
            
            # Question text and marks table
            q_table = doc.add_table(rows=1, cols=2)
            q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            q_table.autofit = False
            
            c_qtext = q_table.cell(0, 0)
            c_qtext.width = Inches(6.3)
            c_qmarks = q_table.cell(0, 1)
            c_qmarks.width = Inches(1.1)
            
            # No borders on question rows
            for c in [c_qtext, c_qmarks]:
                set_cell_border(c)
                
            p_qt = c_qtext.paragraphs[0]
            p_qt.paragraph_format.space_before = Pt(4)
            p_qt.paragraph_format.space_after = Pt(2)
            add_html_formatted_text(p_qt, f"{q_num} {q_text}", default_font="Arial", default_size=Pt(10.5), default_bold=True)
            
            p_qm = c_qmarks.paragraphs[0]
            p_qm.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_qm.paragraph_format.space_before = Pt(4)
            p_qm.paragraph_format.space_after = Pt(2)
            run_qm = p_qm.add_run(f"({q_marks} अंक)")
            run_qm.font.name = "Arial"
            run_qm.font.size = Pt(10.5)
            run_qm.font.bold = True
            
            # Passage if present
            passage = q.get("passage", "")
            if passage:
                p_tbl = doc.add_table(rows=1, cols=1)
                p_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                p_cell = p_tbl.cell(0, 0)
                p_cell.width = Inches(7.4)
                set_cell_background(p_cell, "FCFCFC")
                set_cell_border(p_cell, top="single", bottom="single", left="single", right="single",
                                top_sz="6", bottom_sz="6", left_sz="6", right_sz="6")
                p_pass = p_cell.paragraphs[0]
                if q.get("is_poem", False):
                    p_pass.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p_pass.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    
                add_html_formatted_text(p_pass, passage, default_font="Arial", default_size=Pt(10))
                
            # Sub-questions
            sub_questions = q.get("sub_questions", [])
            if sub_questions:
                for sub in sub_questions:
                    sub_num = sub.get("sub_number", "")
                    sub_text = sub.get("sub_text", "")
                    sub_m = sub.get("marks", 1)
                    
                    sub_tbl = doc.add_table(rows=1, cols=2)
                    sub_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    c_st = sub_tbl.cell(0, 0)
                    c_st.width = Inches(6.3)
                    c_sm = sub_tbl.cell(0, 1)
                    c_sm.width = Inches(1.1)
                    for c in [c_st, c_sm]:
                        set_cell_border(c)
                        
                    p_sub = c_st.paragraphs[0]
                    p_sub.paragraph_format.left_indent = Inches(0.2)
                    p_sub.paragraph_format.space_before = Pt(2)
                    p_sub.paragraph_format.space_after = Pt(2)
                    add_html_formatted_text(p_sub, f"{sub_num} {sub_text}", default_font="Arial", default_size=Pt(10))
                    
                    p_smarks = c_sm.paragraphs[0]
                    p_smarks.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p_smarks.paragraph_format.space_before = Pt(2)
                    p_smarks.paragraph_format.space_after = Pt(2)
                    run_sm = p_smarks.add_run(f"({sub_m} अंक)")
                    run_sm.font.name = "Arial"
                    run_sm.font.size = Pt(10)
                    
                    # Items list
                    for item in sub.get("items", []):
                        p_it = doc.add_paragraph()
                        p_it.paragraph_format.left_indent = Inches(0.4)
                        p_it.paragraph_format.space_before = Pt(0)
                        p_it.paragraph_format.space_after = Pt(1)
                        add_html_formatted_text(p_it, str(item), default_font="Arial", default_size=Pt(9.5))

    # Mandatory Footer Text
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(16)
    p_foot.paragraph_format.space_after = Pt(10)
    run_foot = p_foot.add_run("---------ALL THE BEST!!!----------")
    run_foot.font.name = "Arial"
    run_foot.font.size = Pt(12)
    run_foot.font.bold = True
    
    # Optional Answer Key appended to DOCX
    if include_answer_key and answer_key_data:
        doc.add_page_break()
        p_akh = doc.add_paragraph()
        p_akh.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_akh = p_akh.add_run("आदर्श उत्तरतालिका (Model Answer Key)\n")
        run_akh.font.bold = True
        run_akh.font.size = Pt(14)
        run_akh.font.color.rgb = RGBColor(30, 58, 138)
        
        for ans in answer_key_data.get("answers", []):
            p_ak = doc.add_paragraph()
            p_ak.paragraph_format.space_before = Pt(4)
            p_ak.paragraph_format.space_after = Pt(2)
            add_html_formatted_text(
                p_ak,
                f"{ans.get('section_title', '')} - {ans.get('question_number', '')} ({ans.get('marks', 1)} अंक)\n",
                default_font="Arial",
                default_size=Pt(10),
                default_bold=True
            )
            add_html_formatted_text(
                p_ak,
                f"उत्तर: {ans.get('expected_answer', '')}",
                default_font="Arial",
                default_size=Pt(9.5)
            )
            
            for pt in ans.get("points", []):
                p_pt = doc.add_paragraph()
                p_pt.paragraph_format.left_indent = Inches(0.25)
                p_pt.paragraph_format.space_after = Pt(1)
                add_html_formatted_text(p_pt, f"• {pt}", default_font="Arial", default_size=Pt(9))
                
    doc.save(str(output_path))
    return output_path
